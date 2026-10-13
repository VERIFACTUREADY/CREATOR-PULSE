"""Genera docs/05_ECONOMIA_UNITARIA.xlsx (V2, corrige auditoría C-04).

Cambios frente a V1:
- Un coste desglosado INCOMPLETO ya no se trata como total: la columna pasa a
  «DATOS INCOMPLETOS» y no calcula contribución (antes: IF(SUM>0, SUM, agregado)).
- La columna REAL no admite coste agregado: exige las 5 partidas.
- Nunca etiqueta «rentable»: veredicto «Supera umbral (hipótesis)», «Supera umbral»
  (solo REAL completo), «No supera» o «DATOS INCOMPLETOS».
- Añade devoluciones, coste de contenido/colaboradores por venta, fee real opcional,
  test de precio 24,90/29,90/34,90, capital circulante y panel semanal.
Uso: python3 tools/finance/build_economia.py [salida.xlsx]
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

OUT = sys.argv[1] if len(sys.argv) > 1 else 'docs/05_ECONOMIA_UNITARIA.xlsx'
F = 'Arial'
BLUE = Font(name=F, color='0000FF'); BLACK = Font(name=F); BOLD = Font(name=F, bold=True)
GREEN = Font(name=F, color='008000'); H1 = Font(name=F, bold=True, size=14); HDR = Font(name=F, bold=True, color='FFFFFF')
YEL = PatternFill('solid', fgColor='FFFF00'); NIGHT = PatternFill('solid', fgColor='172B37')
GREY = PatternFill('solid', fgColor='D9D9D9'); REDF = PatternFill('solid', fgColor='F4CCCC')
EUR = '#,##0.00 €;(#,##0.00 €);-'; PCT = '0.0%'

wb = Workbook()

def header(ws, row, titles, widths=None):
    for j, t in enumerate(titles, 1):
        c = ws.cell(row, j, t); c.font = HDR; c.fill = NIGHT; c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 34

def inp(ws, r, c, v, fmt=EUR):
    cell = ws.cell(r, c, v); cell.font = BLUE; cell.fill = YEL; cell.number_format = fmt
    return cell

# ---------------- Leyenda ----------------
L = wb.active; L.title = 'Leyenda'
lines = [
    ('Economía unitaria · Descanso Compartido · V2 (13/10/2026)', H1),
    ('Ninguna cifra de este libro es un coste o precio real confirmado salvo lo que el titular escriba en la columna REAL con facturas o cotizaciones firmadas.', BOLD),
    ('', None),
    ('Cómo usarlo', BOLD),
    ('1. Celdas AMARILLAS con texto AZUL = datos editables. Texto NEGRO = fórmulas (no sobrescribir). VERDE = enlace a otra hoja.', BLACK),
    ('2. Columna REAL (hoja Escenarios): rellenar las 5 partidas del coste (filas 7–11). Si falta UNA, la columna muestra «DATOS INCOMPLETOS» y no calcula margen.', BLACK),
    ('3. Una partida que realmente vale 0 se escribe 0; una celda VACÍA significa «no lo sé» y bloquea el cálculo.', BLACK),
    ('4. Escenarios 1–5 = hipótesis de los informes (coste agregado). Reproducen 0,56 / 7,04 / 11,54 / 11,03 / 7,14 €.', BLACK),
    ('5. Veredictos posibles: «Supera umbral (hipótesis)», «No supera (hipótesis)», «Supera umbral», «No supera», «DATOS INCOMPLETOS». Nunca «rentable».', BLACK),
    ('', None),
    ('Formato esperado de una fila del panel semanal (ejemplo de FORMATO, no son datos): semana 2026-W43 · sesiones 1.000 · sesiones con vista de producto 400 · con ATC 40 · con checkout 20 · pedidos 10 · ventas brutas 299,00 € …', BLACK),
    ('', None),
    ('Fuentes de los supuestos', BOLD),
    ('IVA 21 %: tipo general (AEAT, S25); tipo del SKU pendiente de confirmar por asesoría.', BLACK),
    ('Comisión 2,9 % + 0,30 €: supuesto prudencial de los informes; sustituir por la tarifa real del plan (fila «Comisión real»).', BLACK),
    ('Liquidación Shopify Payments ≥ 3 días hábiles (S27). Capital 50 €: presupuesto declarado por el titular. CAC 10 €: ejemplo del Estudio §9.2.', BLACK),
    ('No incluido: cuota Shopify, apps, gestoría, IRPF y publicidad salvo CAC.', BLACK),
]
for i, (t, f) in enumerate(lines, 1):
    c = L.cell(i, 1, t)
    if f: c.font = f
L.column_dimensions['A'].width = 140

# ---------------- Supuestos ----------------
S = wb.create_sheet('Supuestos')
S['A1'] = 'Supuestos globales'; S['A1'].font = H1
header(S, 3, ['Supuesto', 'Valor', 'Fuente / nota'])
sup = [
    ('IVA aplicado al PVP (hipótesis)', 0.21, PCT, 'S25 tipo general; tipo del SKU pendiente (asesoría).'),                 # B4
    ('Comisión de pago variable (hipótesis)', 0.029, PCT, 'Supuesto de los informes.'),                                     # B5
    ('Comisión de pago fija por pedido (hipótesis)', 0.30, EUR, 'Supuesto de los informes.'),                               # B6
    ('Umbral mínimo de contribución antes de marketing', 7.00, EUR, 'Estudio §9.2 (7–8 €).'),                               # B7
    ('Capital disponible para pedidos y reembolsos', 50, EUR, 'Declarado por el titular.'),                                 # B8
    ('CAC de referencia (hipótesis)', 10, EUR, 'Estudio §9.2.'),                                                            # B9
    ('Días naturales hasta cobrar la liquidación (hipótesis)', 5, '0', 'S27: mínimo 3 días hábiles + banco.'),               # B10
]
for i, (l, v, fmt, n) in enumerate(sup, 4):
    S.cell(i, 1, l).font = BLACK; inp(S, i, 2, v, fmt); S.cell(i, 3, n).font = BLACK
S.column_dimensions['A'].width = 52; S.column_dimensions['B'].width = 14; S.column_dimensions['C'].width = 60

# ---------------- Escenarios ----------------
E = wb.create_sheet('Escenarios')
E['A1'] = 'Contribución por pedido (antes de costes fijos e impuestos directos)'; E['A1'].font = H1
E['A2'] = 'Columnas B–F: hipótesis de los informes. Columna G: REAL (solo datos confirmados).'; E['A2'].font = BOLD
header(E, 3, ['Concepto', '1. Desfavorable', '2. Base', '3. Mejor logística', '4. PVP superior', '5. Ejemplo Maestro', 'REAL (confirmado)'])
rows = {
    4: 'PVP con IVA',
    5: 'Tipo de columna',
    6: 'Coste agregado (solo hipótesis; no se admite en REAL)',
    7: '  · Producto',
    8: '  · Envío a domicilio',
    9: '  · Aranceles / derecho 3 € / importación',
    10: '  · Preparación y embalaje',
    11: '  · Otros costes variables no recuperables',
    12: 'Partidas del desglose rellenadas (de 5)',
    13: 'Estado del coste',
    14: 'Coste completo usado',
    15: 'Reserva incidencias y soporte por pedido',
    16: 'Coste esperado de devoluciones por pedido',
    17: 'Coste de contenido / colaboradores por venta',
    18: 'Comisión real % (vacío = usar supuesto)',
    19: 'PVP sin IVA',
    20: 'Comisión de pago',
    21: 'Contribución por pedido',
    22: 'Veredicto frente al umbral',
    23: 'Contribución tras CAC de referencia',
    24: 'Pedidos financiables con el capital (coste completo)',
}
for r, t in rows.items():
    E.cell(r, 1, t).font = BOLD if r in (13, 14, 21, 22) else BLACK
scen = [(24.90, 16.00, 3.00), (29.90, 14.00, 2.50), (29.90, 10.00, 2.00), (34.90, 14.00, 2.50), (29.90, 14.40, 2.00)]
for j in range(2, 8):
    col = E.cell(1, j).column_letter
    real = (j == 7)
    if real:
        for r in (4, 7, 8, 9, 10, 11, 15, 16, 17, 18): inp(E, r, j, None)
        E.cell(5, j, 'REAL').font = BOLD
        c = E.cell(6, j, 'n/a'); c.fill = GREY; c.font = BLACK
    else:
        pvp, cost, res = scen[j - 2]
        inp(E, 4, j, pvp); E.cell(5, j, 'HIPÓTESIS').font = BLACK
        inp(E, 6, j, cost)
        for r in (7, 8, 9, 10, 11): inp(E, r, j, None)
        inp(E, 15, j, res); inp(E, 16, j, 0); inp(E, 17, j, 0); inp(E, 18, j, None, PCT)
    E.cell(18, j).number_format = PCT
    E.cell(12, j, f'=COUNT({col}7:{col}11)').number_format = '0'
    if real:
        E.cell(13, j, f'=IF(OR({col}4="",{col}12<5,{col}15="",{col}16="",{col}17=""),"DATOS INCOMPLETOS","COMPLETO")')
    else:
        E.cell(13, j, f'=IF({col}12=5,"DESGLOSE COMPLETO",IF(AND({col}12=0,{col}6<>""),"HIPÓTESIS (agregado)","DATOS INCOMPLETOS"))')
    E.cell(14, j, f'=IF({col}13="DATOS INCOMPLETOS","—",IF({col}12=5,SUM({col}7:{col}11),{col}6))').number_format = EUR
    ok = f'{col}13<>"DATOS INCOMPLETOS"'
    E.cell(19, j, f'=IF({col}4="","—",{col}4/(1+Supuestos!$B$4))').number_format = EUR
    E.cell(20, j, f'=IF({col}4="","—",{col}4*IF({col}18="",Supuestos!$B$5,{col}18)+Supuestos!$B$6)').number_format = EUR
    c = E.cell(21, j, f'=IF({ok},{col}19-{col}14-{col}15-{col}16-{col}17-{col}20,"DATOS INCOMPLETOS")'); c.number_format = EUR; c.font = BOLD
    suffix = '' if real else ' (hipótesis)'
    E.cell(22, j, f'=IF({ok},IF({col}21>=Supuestos!$B$7,"Supera umbral{suffix}","No supera{suffix}"),"DATOS INCOMPLETOS")').font = BOLD
    E.cell(23, j, f'=IF({ok},{col}21-Supuestos!$B$9,"DATOS INCOMPLETOS")').number_format = EUR
    E.cell(24, j, f'=IF({ok},INT(Supuestos!$B$8/{col}14),"DATOS INCOMPLETOS")').number_format = '0'
E.conditional_formatting.add('B13:G24', FormulaRule(formula=['B13="DATOS INCOMPLETOS"'], fill=REDF))
notes = [
    'Notas',
    'Fila 13: el fallo de la V1 era IF(SUM(G6:G10)>0,SUM,agregado): con 1 de 5 partidas tomaba ese importe como coste total e inflaba el margen. Ahora 1–4 partidas = DATOS INCOMPLETOS.',
    'Filas 16–17 de las hipótesis valen 0 para reproducir los informes. En REAL deben rellenarse (0 solo si es real).',
    'Ningún veredicto dice «rentable»: superar el umbral antes de marketing no implica beneficio (ver fila 23).',
]
for k, t in enumerate(notes, 26):
    E.cell(k, 1, t).font = BOLD if k == 26 else BLACK
E.column_dimensions['A'].width = 52
for j in range(2, 8): E.column_dimensions[E.cell(1, j).column_letter].width = 18
E.freeze_panes = 'B4'

# ---------------- Precio ----------------
P = wb.create_sheet('Test de precio')
P['A1'] = 'Test de precio con el mismo coste (hipótesis — no son precios decididos)'; P['A1'].font = H1
P['A2'] = 'Coste usado:'; P['B2'] = '=IF(Escenarios!G13="COMPLETO",Escenarios!G14,Escenarios!C14)'; P['B2'].font = GREEN; P['B2'].number_format = EUR
P['C2'] = '=IF(Escenarios!G13="COMPLETO","REAL","HIPÓTESIS: coste del escenario Base")'; P['C2'].font = GREEN
P['A3'] = 'Reserva incidencias:'; P['B3'] = '=IF(Escenarios!G13="COMPLETO",Escenarios!G15,Escenarios!C15)'; P['B3'].font = GREEN; P['B3'].number_format = EUR
header(P, 5, ['PVP con IVA (a probar)', 'PVP sin IVA', 'Comisión', 'Contribución', 'Veredicto', 'Tras CAC ref.'])
for i, pvp in enumerate([24.90, 29.90, 34.90], 6):
    inp(P, i, 1, pvp)
    P.cell(i, 2, f'=A{i}/(1+Supuestos!$B$4)').number_format = EUR
    P.cell(i, 3, f'=A{i}*Supuestos!$B$5+Supuestos!$B$6').number_format = EUR
    P.cell(i, 4, f'=B{i}-$B$2-$B$3-C{i}').number_format = EUR
    P.cell(i, 5, f'=IF(D{i}>=Supuestos!$B$7,"Supera umbral","No supera")&IF($C$2="REAL",""," (hipótesis)")')
    P.cell(i, 6, f'=D{i}-Supuestos!$B$9').number_format = EUR
for c, w in zip('ABCDEF', (24, 14, 12, 14, 30, 14)): P.column_dimensions[c].width = w
P.column_dimensions['C'].width = 40

# ---------------- Caja ----------------
C = wb.create_sheet('Capital circulante')
C['A1'] = 'Capital necesario antes de cobrar las liquidaciones'; C['A1'].font = H1
header(C, 3, ['Concepto', 'Valor', 'Nota'])
C['A4'] = 'Pedidos por día (hipótesis)'; inp(C, 4, 2, 1, '0.0'); C['C4'] = 'Editar: escenario de ventas a cubrir'
C['A5'] = 'Días hasta cobrar'; C['B5'] = '=Supuestos!B10'; C['B5'].font = GREEN
C['A6'] = 'Coste por pedido pagado al proveedor'; C['B6'] = '=IF(Escenarios!G13="COMPLETO",Escenarios!G14,Escenarios!C14)'; C['B6'].font = GREEN; C['B6'].number_format = EUR
C['C6'] = '=IF(Escenarios!G13="COMPLETO","REAL","HIPÓTESIS (Base)")'
C['A7'] = 'Colchón para reembolsos (pedidos equivalentes)'; inp(C, 7, 2, 2, '0'); C['C7'] = 'Hipótesis'
C['A8'] = 'Capital necesario'; C['B8'] = '=B4*B5*B6+B7*B6'; C['B8'].number_format = EUR; C['A8'].font = BOLD; C['B8'].font = BOLD
C['A9'] = 'Capital disponible'; C['B9'] = '=Supuestos!B8'; C['B9'].font = GREEN; C['B9'].number_format = EUR
C['A10'] = '¿Suficiente?'; C['B10'] = '=IF(B9>=B8,"Sí","INSUFICIENTE")'; C['B10'].font = BOLD
C['A11'] = 'Pedidos/día sostenibles sin capital extra'; C['B11'] = '=IF(B5*B6=0,"—",MAX(0,(B9-B7*B6)/(B5*B6)))'; C['B11'].number_format = '0.0'
for c, w in zip('ABC', (48, 16, 40)): C.column_dimensions[c].width = w

# ---------------- Panel semanal ----------------
W = wb.create_sheet('Panel semanal')
W['A1'] = 'Panel semanal (rellenar con datos de Shopify Analytics; sin datos = vacío)'; W['A1'].font = H1
cols = ['Semana', 'Sesiones', 'Sesiones con vista de producto', 'Sesiones con ATC', 'Sesiones con checkout', 'Pedidos', 'Ventas brutas (€)',
        'Gasto de captación (€)', 'Coste de contenido (€)', 'Devoluciones (nº)',
        'Conversión tienda', 'ATC / vistas producto', 'Checkout / ATC', 'AOV', 'CAC', 'Contribución estimada', 'Tasa devolución']
header(W, 3, cols)
for i in range(4, 16):
    for j in range(1, 11):
        inp(W, i, j, None, '#,##0' if j not in (7, 8, 9) else EUR)
    W.cell(i, 1).number_format = '@'
    W.cell(i, 11, f'=IF(OR(B{i}="",B{i}=0,F{i}=""),"",F{i}/B{i})').number_format = PCT
    W.cell(i, 12, f'=IF(OR(C{i}="",C{i}=0,D{i}=""),"",D{i}/C{i})').number_format = PCT
    W.cell(i, 13, f'=IF(OR(D{i}="",D{i}=0,E{i}=""),"",E{i}/D{i})').number_format = PCT
    W.cell(i, 14, f'=IF(OR(F{i}="",F{i}=0,G{i}=""),"",G{i}/F{i})').number_format = EUR
    W.cell(i, 15, f'=IF(OR(F{i}="",F{i}=0,H{i}=""),"",(H{i}+IF(I{i}="",0,I{i}))/F{i})').number_format = EUR
    W.cell(i, 16, f'=IF(F{i}="","",IF(Escenarios!$G$13="COMPLETO",F{i}*Escenarios!$G$21,"DATOS INCOMPLETOS"))').number_format = EUR
    W.cell(i, 17, f'=IF(OR(F{i}="",F{i}=0,J{i}=""),"",J{i}/F{i})').number_format = PCT
W['A17'] = 'Definiciones: conversión = pedidos/sesiones · ATC = sesiones con añadir al carrito / sesiones con vista de producto · Checkout = sesiones con checkout / sesiones con ATC · AOV = ventas brutas / pedidos · CAC = (captación + contenido) / pedidos. Las visualizaciones en redes NO son sesiones.'
for j in range(1, 18): W.column_dimensions[W.cell(1, j).column_letter].width = 14
W.column_dimensions['A'].width = 12

# ---------------- Sensibilidad ----------------
T = wb.create_sheet('Sensibilidad')
T['A1'] = 'Sensibilidad de adquisición (hipótesis de planificación)'; T['A1'].font = H1
T['A2'] = 'Contribución usada:'; T['C2'] = '=IF(Escenarios!G13="COMPLETO",Escenarios!G21,Escenarios!C21)'; T['C2'].font = GREEN; T['C2'].number_format = EUR
T['D2'] = '=IF(Escenarios!G13="COMPLETO","REAL","HIPÓTESIS (Base)")'
header(T, 4, ['Sesiones cualificadas/mes', 'Conversión hipotética', 'Pedidos teóricos', 'Contribución total', 'Capital a adelantar'])
for i, (ses, cv) in enumerate([(300, 0.005), (1000, 0.01), (3000, 0.01), (10000, 0.015)], 5):
    inp(T, i, 1, ses, '#,##0'); inp(T, i, 2, cv, PCT)
    T.cell(i, 3, f'=A{i}*B{i}').number_format = '0.0'
    T.cell(i, 4, f'=C{i}*$C$2').number_format = EUR
    T.cell(i, 5, f'=C{i}*\'Capital circulante\'!$B$6').number_format = EUR
for c, w in zip('ABCDE', (26, 20, 18, 20, 22)): T.column_dimensions[c].width = w

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            f = c.font; c.font = Font(name=F, bold=f.bold, size=f.size, color=f.color)
wb.save(OUT)
print('ok', OUT)
