"""Genera docs/05_ECONOMIA_UNITARIA.xlsx · V2.2 · mercado EE. UU. (USD).

V2.2 (auditoría V2.1, P1):
- Panel semanal con contribución ANTES y DESPUÉS de captación. El contenido ya no se
  cuenta dos veces: «Costes variables» fila 11 = solo comisiones POR VENTA (afiliados,
  creadores); la PRODUCCIÓN de contenido es captación y va solo en el panel (columna I).
- Validación de entradas: reglas de validación de datos en las celdas editables y estado
  «DATOS NO VÁLIDOS» (negativos, porcentajes > 100 %, texto en celdas numéricas, embudo
  incoherente) que bloquea los cálculos aunque se peguen valores saltándose la validación.

Corrige la revisión externa V2:
- «Test de precio» descuenta TODAS las partidas variables (devoluciones esperadas,
  contenido/colaboradores, reserva, atención, envío, aranceles…) y usa la comisión
  real cuando existe (antes omitía G16/G17 y usaba siempre la comisión global).
- Sin IVA español: en EE. UU. el precio se anuncia sin impuestos y el sales tax se
  cobra en el checkout y se ingresa al estado (no es ingreso del vendedor).
- Ningún coste, precio ni umbral real inventado: las celdas de datos están vacías y la
  hoja muestra «DATOS INCOMPLETOS» hasta que el titular las rellene. Los precios de la
  hoja «Test de precio» son PUNTOS DE PRUEBA editables, no precios decididos.
- Sin LTV ni suscripción: el producto es reutilizable (margen del primer pedido).
Uso: python3 tools/finance/build_economia_us.py [salida.xlsx]
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

OUT = sys.argv[1] if len(sys.argv) > 1 else 'docs/05_ECONOMIA_UNITARIA.xlsx'
F = 'Arial'
BLUE = Font(name=F, color='0000FF'); BLACK = Font(name=F); BOLD = Font(name=F, bold=True)
GREEN = Font(name=F, color='008000'); H1 = Font(name=F, bold=True, size=14); HDR = Font(name=F, bold=True, color='FFFFFF')
YEL = PatternFill('solid', fgColor='FFFF00'); NIGHT = PatternFill('solid', fgColor='172B37'); REDF = PatternFill('solid', fgColor='F4CCCC')
USD = '$#,##0.00;($#,##0.00);-'; PCT = '0.0%'
INC = 'DATOS INCOMPLETOS'
INV = 'DATOS NO VÁLIDOS'


def validate(ws, ref, kind):
    """Regla de validación de datos (Excel/LibreOffice). kind: usd (≥ 0), pct (0–1), int (entero ≥ 0), price (> 0), pick (1–3)."""
    cfg = {
        'usd': dict(type='decimal', operator='greaterThanOrEqual', formula1='0', error='Introduce un importe en USD mayor o igual que 0.'),
        'pct': dict(type='decimal', operator='between', formula1='0', formula2='1', error='Introduce un porcentaje entre 0 % y 100 %.'),
        'int': dict(type='whole', operator='greaterThanOrEqual', formula1='0', error='Introduce un número entero mayor o igual que 0.'),
        'price': dict(type='decimal', operator='greaterThan', formula1='0', error='El precio de prueba debe ser mayor que 0.'),
        'pick': dict(type='whole', operator='between', formula1='1', formula2='3', error='Elige 1, 2 o 3.'),
    }[kind]
    dv = DataValidation(allow_blank=True, showErrorMessage=True, errorStyle='stop', errorTitle='Dato no válido', **cfg)
    ws.add_data_validation(dv); dv.add(ref)

wb = Workbook()


def header(ws, row, titles):
    for j, t in enumerate(titles, 1):
        c = ws.cell(row, j, t); c.font = HDR; c.fill = NIGHT; c.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[row].height = 34


def inp(ws, r, c, v, fmt=USD):
    cell = ws.cell(r, c, v); cell.font = BLUE; cell.fill = YEL; cell.number_format = fmt
    return cell


# ------------------------------------------------------------------ Leyenda
L = wb.active; L.title = 'Leyenda'
lines = [
    ('Economía unitaria · SILEXORA · Mercado EE. UU. · USD · V2.2 (14/10/2026)', H1),
    ('Ninguna celda contiene costes, precios ni umbrales reales: se rellenan con cotizaciones, facturas y decisiones del titular. Mientras falte un dato obligatorio, la hoja muestra «DATOS INCOMPLETOS».', BOLD),
    ('', None),
    ('Cómo usarlo', BOLD),
    ('1. Celdas AMARILLAS con texto AZUL = datos editables. NEGRO = fórmulas. VERDE = enlaces entre hojas.', BLACK),
    ('2. Rellena «Costes variables» (columna B) con datos reales por unidad/pedido en USD. Una celda VACÍA = «no lo sé» y bloquea el cálculo; un 0 escrito es un dato válido.', BLACK),
    ('3. «Test de precio»: los precios son PUNTOS DE PRUEBA editables (no precios decididos). Cada fila descuenta TODAS las partidas variables y la comisión real si existe.', BLACK),
    ('4. Impuestos: en EE. UU. el precio se muestra sin sales tax; el impuesto se calcula en el checkout y se ingresa al estado. No se resta del precio (a diferencia del IVA español de la V2).', BLACK),
    ('5. No se modela LTV ni suscripción: el producto es reutilizable. La decisión se basa en el margen del primer pedido.', BLACK),
    ('6. Contenido SIN doble conteo: «Costes variables» fila 11 = solo comisiones que se pagan POR CADA VENTA (afiliados, creadores). La PRODUCCIÓN de contenido (cuotas, muestras, edición) es gasto de captación y se anota solo en «Panel semanal», columna I.', BLACK),
    ('7. «Panel semanal»: contribución ANTES de captación (pedidos × contribución unitaria) y DESPUÉS de captación (− anuncios − producción de contenido). CAC = (anuncios + producción de contenido) / pedidos.', BLACK),
    ('8. Validación: las celdas editables rechazan negativos, porcentajes fuera de 0–100 % y texto. Si se pegan valores no válidos, la hoja muestra «DATOS NO VÁLIDOS» y no calcula.', BLACK),
    ('', None),
    ('Supuestos que hay que verificar', BOLD),
    ('Comisión de pago: 2,9 % + 0,30 $ es un supuesto de partida (tarifa online publicada para planes básicos de Shopify Payments en EE. UU.); confirmar con el plan contratado.', BLACK),
    ('Aranceles: el tratamiento de los envíos de bajo valor importados en EE. UU. (de minimis) ha cambiado recientemente; confirmar con el transitario o agente de aduanas antes de rellenar la partida.', BLACK),
    ('Capital: el titular declaró 1–50 € para España; introducir el importe disponible en USD.', BLACK),
    ('Versión anterior (España, EUR, con IVA): docs/archive/05_ECONOMIA_UNITARIA_ES_V2.xlsx.', BLACK),
]
for i, (t, f) in enumerate(lines, 1):
    c = L.cell(i, 1, t)
    if f: c.font = f
L.column_dimensions['A'].width = 150

# ------------------------------------------------------------------ Supuestos
S = wb.create_sheet('Supuestos')
S['A1'] = 'Supuestos globales (USD)'; S['A1'].font = H1
header(S, 3, ['Supuesto', 'Valor', 'Fuente / nota'])
sup = [
    ('Comisión de pago variable (supuesto)', 0.029, PCT, 'Verificar con el plan de Shopify Payments contratado.'),       # B4
    ('Comisión de pago fija por pedido (supuesto)', 0.30, USD, 'Verificar con el plan contratado.'),                        # B5
    ('Umbral mínimo de contribución por pedido', None, USD, 'Decisión del titular. Vacío = sin veredicto.'),               # B6
    ('Capital disponible para pedidos y reembolsos (USD)', None, USD, 'Dato del titular.'),                                 # B7
    ('CAC de referencia por cliente (USD)', None, USD, 'Dato o hipótesis del titular. Vacío = no se calcula.'),            # B8
    ('Días naturales hasta cobrar la liquidación', None, '0', 'Verificar con Shopify Payments (calendario de pagos).'),   # B9
]
for i, (l, v, fmt, n) in enumerate(sup, 4):
    S.cell(i, 1, l).font = BLACK; inp(S, i, 2, v, fmt); S.cell(i, 3, n).font = BLACK
validate(S, 'B4', 'pct'); validate(S, 'B5:B8', 'usd'); validate(S, 'B9', 'int')
S.column_dimensions['A'].width = 54; S.column_dimensions['B'].width = 14; S.column_dimensions['C'].width = 70

# ------------------------------------------------------------------ Costes variables (REAL)
C = wb.create_sheet('Costes variables')
C['A1'] = 'Costes variables REALES por pedido de 1 unidad (USD)'; C['A1'].font = H1
C['A2'] = 'Solo datos confirmados (cotizaciones, facturas, tarifas). Vacío = bloquea el cálculo.'; C['A2'].font = BOLD
header(C, 3, ['Partida', 'Valor', 'Obligatoria', 'Fuente / nota'])
rows = [
    ('Producto: coste unitario del proveedor', USD, 'Sí', 'Cotización firmada del modelo exacto.'),                         # 4
    ('Flete internacional / inbound por unidad', USD, 'Sí', 'Prorrateado por unidad.'),                                    # 5
    ('Aranceles, tasas y agente de aduanas por unidad', USD, 'Sí', 'Confirmar régimen vigente con el agente de aduanas.'), # 6
    ('Envío al cliente (última milla)', USD, 'Sí', 'Tarifa real del transportista.'),                                       # 7
    ('Embalaje y preparación (fulfillment)', USD, 'Sí', ''),                                                                 # 8
    ('Reserva por incidencias / contracargos por pedido', USD, 'Sí', 'Pérdidas, daños, fraude.'),                           # 9
    ('Atención al cliente por pedido', USD, 'Sí', 'Tiempo o herramientas imputables.'),                                     # 10
    ('Comisiones POR VENTA a afiliados / creadores', USD, 'Sí', 'Solo lo que se paga por cada venta. La PRODUCCIÓN de contenido NO va aquí (es captación: Panel semanal, col. I). 0 si no hay.'),  # 11
    ('Tasa de devolución esperada', PCT, 'Sí', 'Dato real o prudente; 0 % solo si se documenta.'),                          # 12
    ('Coste por devolución (envío de vuelta + producto no recuperable)', USD, 'Sí', 'Por cada devolución.'),               # 13
    ('Ingreso por envío cobrado al cliente', USD, 'Sí', '0 si el envío es gratuito para el cliente.'),                     # 14
    ('Descuento medio por pedido', USD, 'Sí', '0 si no hay descuentos.'),                                                   # 15
    ('Comisión de pago variable REAL (opcional)', PCT, 'No', 'Vacío = usar Supuestos!B4.'),                                 # 16
    ('Comisión de pago fija REAL (opcional)', USD, 'No', 'Vacío = usar Supuestos!B5.'),                                     # 17
]
for i, (l, fmt, req, n) in enumerate(rows, 4):
    C.cell(i, 1, l).font = BLACK; inp(C, i, 2, None, fmt); C.cell(i, 3, req).font = BLACK; C.cell(i, 4, n).font = BLACK
    validate(C, f'B{i}', 'pct' if fmt == PCT else 'usd')
C['A19'] = 'Partidas obligatorias rellenadas (de 12)'; C['B19'] = '=COUNT(B4:B15)'
C['A28'] = 'Entradas no válidas (negativos, % fuera de 0–100 %, texto), aquí y en «Supuestos»'
C['B28'] = ('=COUNTIF(B4:B17,"<0")+(COUNTA(B4:B17)-COUNT(B4:B17))+IF(ISNUMBER(B12),--(B12>1),0)+IF(ISNUMBER(B16),--(B16>1),0)'
            '+COUNTIF(Supuestos!B4:B9,"<0")+(COUNTA(Supuestos!B4:B9)-COUNT(Supuestos!B4:B9))+IF(ISNUMBER(Supuestos!B4),--(Supuestos!B4>1),0)')
C['A20'] = 'Estado'; C['B20'] = f'=IF(B28>0,"{INV}",IF(B19=12,"COMPLETO","{INC}"))'; C['A20'].font = BOLD; C['B20'].font = BOLD
C['A21'] = 'Coste de producto puesto en cliente (filas 4–8)'; C['B21'] = f'=IF($B$20="COMPLETO",SUM(B4:B8),$B$20)'
C['A22'] = 'Coste esperado de devoluciones por pedido (fila 12 × fila 13)'; C['B22'] = f'=IF($B$20="COMPLETO",B12*B13,$B$20)'
C['A23'] = 'Otros costes variables por pedido (filas 9–11: incidencias, atención, comisiones por venta)'; C['B23'] = f'=IF($B$20="COMPLETO",SUM(B9:B11),$B$20)'
C['A24'] = 'Coste variable total por pedido (antes de comisión)'; C['B24'] = f'=IF($B$20="COMPLETO",B21+B22+B23,$B$20)'
C['A25'] = 'Comisión variable usada'; C['B25'] = '=IF(B16="",Supuestos!B4,B16)'  # solo se usa si B20 = COMPLETO (que ya exige entradas válidas)
C['A26'] = 'Comisión fija usada'; C['B26'] = '=IF(B17="",Supuestos!B5,B17)'
C['A27'] = 'Desembolso antes de cobrar (filas 4–8)'; C['B27'] = f'=IF($B$20="COMPLETO",SUM(B4:B8),$B$20)'
for r in (21, 22, 23, 24, 26, 27): C.cell(r, 2).number_format = USD
C['B25'].number_format = PCT
for r in (24,): C.cell(r, 1).font = BOLD; C.cell(r, 2).font = BOLD
C.conditional_formatting.add('B20:B28', FormulaRule(formula=[f'OR($B$20="{INC}",$B$20="{INV}")'], fill=REDF))
C.column_dimensions['A'].width = 62; C.column_dimensions['B'].width = 16; C.column_dimensions['C'].width = 12; C.column_dimensions['D'].width = 54

# ------------------------------------------------------------------ Test de precio
P = wb.create_sheet('Test de precio')
P['A1'] = 'Test de precio · puntos de prueba en USD (no son precios decididos)'; P['A1'].font = H1
P['A2'] = 'Cada fila descuenta TODAS las partidas de «Costes variables» y la comisión real si existe.'; P['A2'].font = BOLD
cols = ['Precio de prueba (sin sales tax)', 'Ingreso del pedido (precio − descuento + envío cobrado)', 'Comisión de pago', 'Coste producto puesto en cliente',
        'Devoluciones esperadas', 'Incidencias + atención + comisiones por venta', 'Contribución por pedido ANTES de captación', 'Veredicto frente al umbral', 'Contribución DESPUÉS de captación (CAC de referencia)']
header(P, 4, cols)
for i, price in enumerate([24.99, 29.99, 34.99], 5):
    inp(P, i, 1, price)
    ok = "'Costes variables'!$B$20=\"COMPLETO\""
    st = "'Costes variables'!$B$20"
    P.cell(i, 2, f'=IF(A{i}="","—",A{i}-\'Costes variables\'!$B$15+\'Costes variables\'!$B$14)').number_format = USD
    P.cell(i, 3, f'=IF(A{i}="","—",B{i}*\'Costes variables\'!$B$25+\'Costes variables\'!$B$26)').number_format = USD
    P.cell(i, 4, "='Costes variables'!$B$21").number_format = USD
    P.cell(i, 5, "='Costes variables'!$B$22").number_format = USD
    P.cell(i, 6, "='Costes variables'!$B$23").number_format = USD
    c = P.cell(i, 7, f'=IF(A{i}="","{INC}",IF(OR(NOT(ISNUMBER(A{i})),{st}="{INV}"),"{INV}",IF(A{i}<=0,"{INV}",IF(NOT({ok}),"{INC}",B{i}-C{i}-D{i}-E{i}-F{i}))))'); c.number_format = USD; c.font = BOLD
    validate(P, f'A{i}', 'price')
    P.cell(i, 8, f'=IF(NOT(ISNUMBER(G{i})),G{i},IF(Supuestos!$B$6="","Umbral sin definir",IF(G{i}>=Supuestos!$B$6,"Supera umbral","No supera")))').font = BOLD
    P.cell(i, 9, f'=IF(NOT(ISNUMBER(G{i})),G{i},IF(Supuestos!$B$8="","CAC sin definir",G{i}-Supuestos!$B$8))').number_format = USD
    for j in (4, 5, 6): P.cell(i, j).font = GREEN
P.conditional_formatting.add('G5:I7', FormulaRule(formula=['NOT(ISNUMBER($G5))'], fill=REDF))
P['A9'] = 'Comprobación: contribución antes de captación = ingreso − comisión − coste producto − devoluciones − (incidencias + atención + comisiones por venta). Después de captación = antes − CAC de referencia (Supuestos!B8). Ninguna partida se omite ni se cuenta dos veces.'
for j, w in enumerate([18, 24, 14, 18, 16, 20, 18, 22, 18], 1): P.column_dimensions[P.cell(1, j).column_letter].width = w

# ------------------------------------------------------------------ Capital circulante
K = wb.create_sheet('Capital circulante')
K['A1'] = 'Capital necesario antes de cobrar las liquidaciones (USD)'; K['A1'].font = H1
header(K, 3, ['Concepto', 'Valor', 'Nota'])
K['A4'] = 'Pedidos por día a cubrir'; inp(K, 4, 2, None, '0.0'); K['C4'] = 'Escenario de ventas del titular'; validate(K, 'B4', 'usd')
K['A5'] = 'Días hasta cobrar'; K['B5'] = '=Supuestos!B9'; K['B5'].font = GREEN
K['A6'] = 'Desembolso por pedido antes de cobrar'; K['B6'] = "='Costes variables'!B27"; K['B6'].font = GREEN; K['B6'].number_format = USD
K['A7'] = 'Colchón para reembolsos (pedidos equivalentes)'; inp(K, 7, 2, None, '0'); K['C7'] = 'Decisión del titular'; validate(K, 'B7', 'int')
K['A8'] = 'Capital necesario'; K['B8'] = f'=IF(OR(B4="",B5="",B7=""),"{INC}",IF(OR(NOT(ISNUMBER(B4)),NOT(ISNUMBER(B7)),B4<0,B7<0),"{INV}",IF(NOT(ISNUMBER(B6)),B6,B4*B5*B6+B7*B6)))'; K['B8'].number_format = USD; K['A8'].font = BOLD; K['B8'].font = BOLD
K['A9'] = 'Capital disponible'; K['B9'] = '=IF(Supuestos!B7="","—",Supuestos!B7)'; K['B9'].font = GREEN; K['B9'].number_format = USD
K['A10'] = '¿Suficiente?'; K['B10'] = f'=IF(NOT(ISNUMBER(B8)),B8,IF(B9="—","{INC}",IF(B9>=B8,"Sí","INSUFICIENTE")))'; K['B10'].font = BOLD
for c, w in zip('ABC', (48, 16, 40)): K.column_dimensions[c].width = w

# ------------------------------------------------------------------ Panel semanal
W = wb.create_sheet('Panel semanal')
W['A1'] = 'Panel semanal (datos de Shopify Analytics, en USD)'; W['A1'].font = H1
W['A2'] = 'Precio de venta actual (USD, sin sales tax):'; inp(W, 2, 4, None); validate(W, 'D2', 'price')
cols = ['Semana', 'Sesiones', 'Sesiones con vista de producto', 'Sesiones con ATC', 'Sesiones con checkout', 'Pedidos', 'Ventas brutas (USD)',
        'Gasto en anuncios (USD)', 'Producción de contenido (USD; sin comisiones por venta)', 'Devoluciones (nº)',
        'Conversión tienda', 'ATC / vistas producto', 'Checkout / ATC', 'AOV', 'CAC (anuncios + contenido) / pedidos',
        'Contribución ANTES de captación', 'Captación total (H + I)', 'Contribución DESPUÉS de captación', 'Tasa devolución', 'Validación de la fila']
header(W, 3, cols)
# Contribución unitaria: ya incluye las comisiones POR VENTA (Costes variables, fila 11). La producción de contenido (col. I) NO.
cu = (f"IF($D$2=\"\",\"{INC}\",IF(OR(NOT(ISNUMBER($D$2)),'Costes variables'!$B$20=\"{INV}\"),\"{INV}\",IF($D$2<=0,\"{INV}\","
      f"IF('Costes variables'!$B$20<>\"COMPLETO\",\"{INC}\","
      f"($D$2-'Costes variables'!$B$15+'Costes variables'!$B$14)*(1-'Costes variables'!$B$25)-'Costes variables'!$B$26-'Costes variables'!$B$24))))")
W['A19'] = 'Contribución unitaria usada en el panel:'; W['D19'] = '=' + cu; W['D19'].number_format = USD; W['D19'].font = BOLD
for i in range(4, 16):
    for j in range(1, 11):
        inp(W, i, j, None, '#,##0' if j not in (7, 8, 9) else USD)
    W.cell(i, 1).number_format = '@'
    T = f'T{i}'
    W[T] = (f'=IF(COUNTA(B{i}:J{i})=0,"",IF(OR(COUNTA(B{i}:J{i})<>COUNT(B{i}:J{i}),COUNTIF(B{i}:J{i},"<0")>0,'
            f'AND(ISNUMBER(C{i}),ISNUMBER(B{i}),C{i}>B{i}),AND(ISNUMBER(D{i}),ISNUMBER(C{i}),D{i}>C{i}),'
            f'AND(ISNUMBER(E{i}),ISNUMBER(D{i}),E{i}>D{i})),"{INV}","OK"))')
    W[T].font = BOLD
    g = lambda body: f'=IF(${T}="{INV}","{INV}",{body})'
    W.cell(i, 11, g(f'IF(OR(B{i}="",B{i}=0,F{i}=""),"",F{i}/B{i})')).number_format = PCT
    W.cell(i, 12, g(f'IF(OR(C{i}="",C{i}=0,D{i}=""),"",D{i}/C{i})')).number_format = PCT
    W.cell(i, 13, g(f'IF(OR(D{i}="",D{i}=0,E{i}=""),"",E{i}/D{i})')).number_format = PCT
    W.cell(i, 14, g(f'IF(OR(F{i}="",F{i}=0,G{i}=""),"",G{i}/F{i})')).number_format = USD
    W.cell(i, 15, g(f'IF(OR(F{i}="",F{i}=0,Q{i}=""),"",Q{i}/F{i})')).number_format = USD
    W.cell(i, 16, g(f'IF(F{i}="","",IF(ISNUMBER($D$19),F{i}*$D$19,$D$19))')).number_format = USD
    W.cell(i, 17, g(f'IF(H{i}="","",H{i}+IF(I{i}="",0,I{i}))')).number_format = USD
    W.cell(i, 18, g(f'IF(OR(P{i}="",Q{i}=""),"",IF(ISNUMBER(P{i}),P{i}-Q{i},P{i}))')).number_format = USD
    W.cell(i, 19, g(f'IF(OR(F{i}="",F{i}=0,J{i}=""),"",J{i}/F{i})')).number_format = PCT
validate(W, 'B4:F15', 'int'); validate(W, 'G4:I15', 'usd'); validate(W, 'J4:J15', 'int')
W.conditional_formatting.add('K4:T15', FormulaRule(formula=[f'$T4="{INV}"'], fill=REDF))
W['A17'] = ('Definiciones: conversión = pedidos/sesiones · ATC = sesiones con añadir al carrito / sesiones con vista de producto · checkout = sesiones con checkout / sesiones con ATC · '
            'AOV = ventas brutas / pedidos · CAC = (anuncios + producción de contenido) / pedidos · contribución ANTES de captación = pedidos × contribución unitaria (ya descuenta las comisiones POR VENTA de «Costes variables» fila 11) · '
            'DESPUÉS de captación = ANTES − anuncios − producción de contenido. Cada coste de contenido se cuenta una sola vez. Las visualizaciones en redes no son sesiones.')
W['A18'] = 'Validación: negativos, texto en celdas numéricas o un embudo incoherente (vistas > sesiones, ATC > vistas, checkout > ATC) marcan la fila como «DATOS NO VÁLIDOS» y no se calcula.'
for j in range(1, 21): W.column_dimensions[W.cell(1, j).column_letter].width = 14

# ------------------------------------------------------------------ Sensibilidad
T = wb.create_sheet('Sensibilidad')
T['A1'] = 'Sensibilidad de adquisición (planificación, no previsión)'; T['A1'].font = H1
T['A2'] = 'Contribución usada (Test de precio, fila elegida):'; inp(T, 2, 5, 2, '0'); validate(T, 'E2', 'pick'); T['F2'] = '← 1, 2 o 3 = fila 5, 6 o 7 de «Test de precio»'
T['C3'] = "=INDEX('Test de precio'!G5:G7,E2)"; T['A3'] = 'Contribución por pedido:'; T['C3'].font = GREEN; T['C3'].number_format = USD
header(T, 5, ['Sesiones cualificadas/mes', 'Conversión hipotética', 'Pedidos teóricos', 'Contribución total', 'Capital a adelantar'])
for i in range(6, 10):
    inp(T, i, 1, None, '#,##0'); inp(T, i, 2, None, PCT); validate(T, f'A{i}', 'int'); validate(T, f'B{i}', 'pct')
    T.cell(i, 3, f'=IF(OR(A{i}="",B{i}=""),"",A{i}*B{i})').number_format = '0.0'
    T.cell(i, 4, f'=IF(C{i}="","",IF(ISNUMBER($C$3),C{i}*$C$3,"{INC}"))').number_format = USD
    T.cell(i, 5, f"=IF(C{i}=\"\",\"\",IF(ISNUMBER('Costes variables'!$B$27),C{i}*'Costes variables'!$B$27,\"{INC}\"))").number_format = USD
for c, w in zip('ABCDEF', (26, 20, 18, 20, 22, 40)): T.column_dimensions[c].width = w

for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            f = c.font; c.font = Font(name=F, bold=f.bold, size=f.size, color=f.color)
wb.save(OUT)
print('ok', OUT)
