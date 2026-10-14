"""Pruebas de regresión de docs/05_ECONOMIA_UNITARIA.xlsx (V2.2, EE. UU., USD).
V2.2: panel con contribución antes/después de captación sin doble conteo de contenido y validación de entradas.
Cubren el hallazgo P1 de la revisión V2: «Test de precio» debe descontar TODAS las partidas
variables (devoluciones, contenido…) y usar la comisión real cuando exista.
Requiere LibreOffice (soffice). Uso: python3 tools/qa/excel_tests.py → qa/reports/excel-tests.json
Los valores numéricos de los casos son DATOS DE PRUEBA, no costes reales."""
import json, os, shutil, subprocess, sys, tempfile
from openpyxl import load_workbook

REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), '../..'))
SRC = os.path.join(REPO, 'docs/05_ECONOMIA_UNITARIA.xlsx')
INC = 'DATOS INCOMPLETOS'
INV = 'DATOS NO VÁLIDOS'

def recalc(path):
    d = tempfile.mkdtemp()
    subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', d, path],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    return load_workbook(os.path.join(d, os.path.basename(path)), data_only=True)

def case(edits):
    d = tempfile.mkdtemp(); p = os.path.join(d, 'caso.xlsx'); shutil.copy(SRC, p)
    wb = load_workbook(p)
    for (sheet, ref), v in edits.items(): wb[sheet][ref] = v
    wb.save(p)
    return recalc(p)

# Costes de PRUEBA (12 obligatorias): filas 4–15 de «Costes variables»
BASE = {('Costes variables', f'B{r}'): v for r, v in zip(range(4, 16), [8, 1, 1, 4, 0.5, 0.5, 0.5, 0, 0.05, 6, 0, 0])}

def expected(price, c, fee=0.029, fixed=0.30):
    product, inbound, duties, lastmile, pack, reserve, support, content, rrate, rcost, shiprev, disc = c
    rev = price - disc + shiprev
    return rev - (rev * fee + fixed) - (product + inbound + duties + lastmile + pack) - rrate * rcost - (reserve + support + content)

res = []
def check(name, ok, detail):
    res.append({'prueba': name, 'ok': bool(ok), 'detalle': detail}); print(('OK  ' if ok else 'FAIL') + ' ' + name + ' → ' + str(detail))

vals = [8, 1, 1, 4, 0.5, 0.5, 0.5, 0, 0.05, 6, 0, 0]
# 1. Entregado vacío: nada se calcula
wb = recalc(SRC); P = wb['Test de precio']
check('Hoja entregada sin datos → DATOS INCOMPLETOS en las 3 filas', all(P[f'G{r}'].value == INC for r in (5, 6, 7)), [P[f'G{r}'].value for r in (5, 6, 7)])
check('Sin IVA español: ninguna fórmula divide por 1,21', not any('1.21' in str(c.value) for ws in load_workbook(SRC).worksheets for row in ws.iter_rows() for c in row), 'ok')
check('Ningún texto «rentable»', not any('rentable' in str(c.value).lower() for ws in load_workbook(SRC).worksheets for row in ws.iter_rows() for c in row if c.data_type == 'f'), 'ok')

# 2. 11 de 12 partidas
e = dict(BASE); e[('Costes variables', 'B11')] = None
P = case(e)['Test de precio']
check('11/12 partidas → DATOS INCOMPLETOS', P['G6'].value == INC, P['G6'].value)

# 3. Completo: la contribución incluye TODAS las partidas
P = case(BASE)['Test de precio']
exp = [round(expected(p, vals), 4) for p in (24.99, 29.99, 34.99)]
got = [P[f'G{r}'].value for r in (5, 6, 7)]
check('Completo: contribución = ingreso − comisión − todas las partidas', [round(x, 4) for x in got] == exp, {'esperado': exp, 'hoja': [round(x, 4) for x in got]})

# 4. Devoluciones > 0: reduce exactamente tasa × coste por devolución
e = dict(BASE); e[('Costes variables', 'B12')] = 0.10
g2 = case(e)['Test de precio']['G6'].value
check('Subir devoluciones 5 %→10 % (coste 6 $) reduce exactamente 0,30 $', abs((got[1] - g2) - 0.30) < 1e-9, round(got[1] - g2, 4))

# 5. Contenido > 0: reduce exactamente su importe
e = dict(BASE); e[('Costes variables', 'B11')] = 2
g3 = case(e)['Test de precio']['G6'].value
check('Contenido 0→2 $ reduce exactamente 2,00 $ [fallo V2]', abs((got[1] - g3) - 2) < 1e-9, round(got[1] - g3, 4))

# 6. Comisión real
e = dict(BASE); e[('Costes variables', 'B16')] = 0.035; e[('Costes variables', 'B17')] = 0.49
g4 = case(e)['Test de precio']['G6'].value
check('Comisión real 3,5 % + 0,49 $ sustituye al supuesto [fallo V2]', abs(g4 - expected(29.99, vals, 0.035, 0.49)) < 1e-9, round(g4, 4))

# 7. Descuento y envío cobrado
v2 = list(vals); v2[10] = 5; v2[11] = 3
e = dict(BASE); e[('Costes variables', 'B14')] = 5; e[('Costes variables', 'B15')] = 3
g5 = case(e)['Test de precio']['G6'].value
check('Envío cobrado 5 $ y descuento 3 $ entran en el ingreso y la comisión', abs(g5 - expected(29.99, v2)) < 1e-9, round(g5, 4))

# 8. Umbral y CAC
wb = case({**BASE, ('Supuestos', 'B6'): None})
check('Umbral vacío → «Umbral sin definir»', wb['Test de precio']['H6'].value == 'Umbral sin definir', wb['Test de precio']['H6'].value)
wb = case({**BASE, ('Supuestos', 'B6'): 7, ('Supuestos', 'B8'): 10})
P = wb['Test de precio']
check('Umbral 7 $ y CAC 10 $ → veredicto y contribución tras CAC', P['H6'].value in ('Supera umbral', 'No supera') and abs(P['I6'].value - (got[1] - 10)) < 1e-9, (P['H6'].value, round(P['I6'].value, 4)))

# 9. Panel semanal coherente con el test de precio
wb = case({**BASE, ('Panel semanal', 'D2'): 29.99, ('Panel semanal', 'F4'): 10, ('Panel semanal', 'B4'): 1000})
p16 = wb['Panel semanal']['P4'].value
check('Panel: contribución estimada = pedidos × contribución unitaria', abs(p16 - 10 * got[1]) < 1e-6, round(p16, 4))

# 10. Capital circulante
wb = case({**BASE, ('Supuestos', 'B7'): 50, ('Supuestos', 'B9'): 5, ('Capital circulante', 'B4'): 1, ('Capital circulante', 'B7'): 2})
K = wb['Capital circulante']
check('Capital: 1 pedido/día × 5 días × 14,50 $ + 2 × 14,50 $ = 101,50 $ → INSUFICIENTE con 50 $', abs(K['B8'].value - 101.5) < 1e-9 and K['B10'].value == 'INSUFICIENTE', (K['B8'].value, K['B10'].value))

# ---------- V2.2 (auditoría V2.1 P1) ----------
# 11. Panel: contribución ANTES y DESPUÉS de captación; contenido contado UNA vez
PANEL = {**BASE, ('Panel semanal', 'D2'): 29.99, ('Panel semanal', 'B4'): 1000, ('Panel semanal', 'C4'): 400, ('Panel semanal', 'D4'): 60,
         ('Panel semanal', 'E4'): 30, ('Panel semanal', 'F4'): 10, ('Panel semanal', 'H4'): 40, ('Panel semanal', 'I4'): 50}
W = case(PANEL)['Panel semanal']
before, total, after, cac = W['P4'].value, W['Q4'].value, W['R4'].value, W['O4'].value
check('Panel V2.2: ANTES = pedidos × contribución; captación = anuncios + contenido; DESPUÉS = ANTES − captación; CAC = captación / pedidos',
      abs(before - 10 * got[1]) < 1e-6 and abs(total - 90) < 1e-9 and abs(after - (10 * got[1] - 90)) < 1e-6 and abs(cac - 9) < 1e-9,
      {'antes': round(before, 4), 'captacion': total, 'despues': round(after, 4), 'cac': cac})
# Sin doble conteo: la producción de contenido (I) NO cambia la contribución ANTES; solo la de DESPUÉS, exactamente una vez
W2 = case({**PANEL, ('Panel semanal', 'I4'): 150})['Panel semanal']
check('Contenido +100 $ en el panel: ANTES igual y DESPUÉS baja exactamente 100 $ (sin doble conteo)',
      abs(W2['P4'].value - before) < 1e-9 and abs((after - W2['R4'].value) - 100) < 1e-6, round(after - W2['R4'].value, 4))
# Comisiones POR VENTA (fila 11) solo en ANTES: 2 $ × 10 pedidos = 20 $ (y no tocan la captación)
W3 = case({**PANEL, ('Costes variables', 'B11'): 2})['Panel semanal']
check('Comisión por venta 2 $ (fila 11): ANTES baja 20 $ con 10 pedidos y la captación no cambia',
      abs((before - W3['P4'].value) - 20) < 1e-6 and W3['Q4'].value == total, round(before - W3['P4'].value, 4))

# 12. Validación de entradas
for label, edit in [('coste negativo (−1 $ en producto)', {('Costes variables', 'B4'): -1}),
                    ('tasa de devolución 500 %', {('Costes variables', 'B12'): 5}),
                    ('texto en una celda numérica («ocho»)', {('Costes variables', 'B4'): 'ocho'}),
                    ('comisión supuesta 290 % (2,9 en lugar de 0,029)', {('Supuestos', 'B4'): 2.9}),
                    ('días hasta cobrar negativos', {('Supuestos', 'B9'): -3})]:
    wb = case({**BASE, **edit})
    check(f'Validación: {label} → «{INV}» y sin contribución', wb['Costes variables']['B20'].value == INV and all(wb['Test de precio'][f'G{r}'].value == INV for r in (5, 6, 7)),
          (wb['Costes variables']['B20'].value, wb['Test de precio']['G6'].value))
P = case({**BASE, ('Test de precio', 'A6'): -29.99})['Test de precio']
check('Validación: precio de prueba negativo → «DATOS NO VÁLIDOS» solo en esa fila', P['G6'].value == INV and isinstance(P['G5'].value, float), (P['G5'].value, P['G6'].value))
W = case({**PANEL, ('Panel semanal', 'D4'): 500})['Panel semanal']
check('Validación: embudo incoherente (ATC 500 > vistas 400) → fila «DATOS NO VÁLIDOS»', W['T4'].value == INV and W['P4'].value == INV and W['R4'].value == INV, (W['T4'].value, W['P4'].value))
W = case({**PANEL, ('Panel semanal', 'F4'): -10})['Panel semanal']
check('Validación: pedidos negativos → fila «DATOS NO VÁLIDOS»', W['T4'].value == INV and W['O4'].value == INV, (W['T4'].value, W['O4'].value))
W = case({**PANEL, ('Panel semanal', 'D2'): -5})['Panel semanal']
check('Validación: precio actual negativo en el panel → contribución «DATOS NO VÁLIDOS»', W['P4'].value == INV, W['P4'].value)
W = case(PANEL)['Panel semanal']
check('Panel con datos válidos → fila «OK»', W['T4'].value == 'OK', W['T4'].value)
wbf = load_workbook(SRC)
dv_refs = {ws.title: ' '.join(str(dv.sqref) for dv in ws.data_validations.dataValidation) for ws in wbf.worksheets}
need = {'Supuestos': ['B4', 'B9'], 'Costes variables': ['B4', 'B12', 'B17'], 'Test de precio': ['A5', 'A7'], 'Panel semanal': ['D2', 'B4:F15', 'G4:I15'], 'Capital circulante': ['B4', 'B7'], 'Sensibilidad': ['E2']}
missing = [f'{sh}!{r}' for sh, refs in need.items() for r in refs if r not in dv_refs.get(sh, '')]
check('Reglas de validación de datos en las celdas editables (Excel/LibreOffice)', not missing, missing or 'ok')

os.makedirs(os.path.join(REPO, 'qa/reports'), exist_ok=True)
json.dump(res, open(os.path.join(REPO, 'qa/reports/excel-tests.json'), 'w'), indent=2, ensure_ascii=False)
print(f"\n{sum(r['ok'] for r in res)}/{len(res)} pruebas OK")
sys.exit(0 if all(r['ok'] for r in res) else 1)
