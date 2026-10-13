"""Demuestra el fallo P1 de la revisión V2 sobre la hoja ARCHIVADA (España, V2):
«Test de precio» no cambiaba al subir devoluciones (G16), contenido (G17) ni la comisión real (G18).
Uso: python3 tools/qa/v2-regression-excel.py → qa/reports/v2-regression-excel.json"""
import json, os, shutil, subprocess, tempfile
from openpyxl import load_workbook
REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), '../..'))
SRC = os.path.join(REPO, 'docs/archive/05_ECONOMIA_UNITARIA_ES_V2.xlsx')
def run(edits):
    d = tempfile.mkdtemp(); p = os.path.join(d, 'c.xlsx'); shutil.copy(SRC, p)
    wb = load_workbook(p); ws = wb['Escenarios']
    for k, v in edits.items(): ws[k] = v
    wb.save(p)
    subprocess.run(['soffice', '--headless', '--calc', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', d + '/o', p], check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    return load_workbook(d + '/o/c.xlsx', data_only=True)['Test de precio']['D7'].value
base = {'G4': 29.90, 'G7': 8, 'G8': 4, 'G9': 1, 'G10': 0.5, 'G11': 0.5, 'G15': 2.5, 'G16': 0, 'G17': 0}
a = run(base); b = run({**base, 'G16': 1.5, 'G17': 2}); c = run({**base, 'G18': 0.05})
out = {'contribucion_base': a, 'con_devoluciones_1_5_y_contenido_2': b, 'con_comision_real_5pct': c,
       'fallo_reproducido': a == b == c}
json.dump(out, open(os.path.join(REPO, 'qa/reports/v2-regression-excel.json'), 'w'), indent=2)
print(out)
